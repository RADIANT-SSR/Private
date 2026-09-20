"""Pinning tests for CU-374 — export formats (GUI usability audit 2026-09).

Each test drives the real export surface (the File-menu action or the exporter it
calls) on the shipped example and asserts the fixed format. Every test here failed on
the pre-fix code; the finding numbers refer to ``docs/reports/gui_usability_audit_2026-09/``.
"""

from __future__ import annotations

import json
import warnings
from pathlib import Path

import pytest

pytest.importorskip("PySide6", reason="GUI tests require the optional 'gui' extra")

from radiant.api.sensor import Sensor  # noqa: E402
from radiant.gui.main_window import RADIANTMainWindow  # noqa: E402

_EXAMPLE = Path(__file__).resolve().parents[4] / "examples" / "mwir_leo_minimal.yaml"
_WAIT_MS = 15000


def _window(qtbot) -> RADIANTMainWindow:  # type: ignore[no-untyped-def]
    window = RADIANTMainWindow(Sensor.load(_EXAMPLE), path=str(_EXAMPLE))
    qtbot.addWidget(window)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        with qtbot.waitSignal(window.evaluationFinished, timeout=_WAIT_MS):
            pass
    return window


def _save_to(monkeypatch, dest: Path) -> None:  # type: ignore[no-untyped-def]
    from radiant.gui import main_window as mw

    monkeypatch.setattr(
        mw.QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(dest), ""))
    )


class TestF42AuditTrail:
    """F-42: Export Resolved YAML carried no per-value provenance, and the JSON
    record's git_commit was `unknown` on a source checkout whose title bar showed it."""

    def test_resolved_yaml_marks_provenance(self, qtbot, tmp_path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
        """J-6.1: Karen's 'exact inputs used for this prediction'."""
        window = _window(qtbot)
        window.sensor.set("readout.read_noise_e_rms", 40.0)
        dest = tmp_path / "resolved.yaml"
        _save_to(monkeypatch, dest)
        window.action("file.export_resolved_yaml").trigger()
        lines = dest.read_text(encoding="utf-8").splitlines()
        assert any(
            line.startswith("  read_noise_e_rms: 40.0") and line.endswith("# user-set")
            for line in lines
        )
        assert any(line.endswith("# default") for line in lines)
        assert any(line.endswith("# derived") for line in lines)

    def test_json_record_carries_the_checkout_commit(self, qtbot, tmp_path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
        """J-7.2: launched from an unrelated directory, the record still names the commit."""
        from radiant.api.build_info import build_info

        expected = build_info().git_sha
        if expected is None:
            pytest.skip("not a git checkout — the record correctly reports unknown")
        monkeypatch.chdir(tmp_path)
        window = _window(qtbot)
        dest = tmp_path / "result.json"
        _save_to(monkeypatch, dest)
        window.action("file.export_json").trigger()
        record = json.loads(dest.read_text(encoding="utf-8"))
        assert record["git_commit"] != "unknown"
        assert record["git_commit"].startswith(expected[:7]) or expected.startswith(
            record["git_commit"][:7]
        )


class TestF17SweepCsvUnitsAndNumbers:
    """F-17 / F-31: the sweep CSV had bare column names, `np.float64(…)` literals in
    twelve cells, float-noise axis values, and codes/flags that read as values."""

    def test_sweep_csv_has_units_plain_numbers_and_typed_axis(
        self, qtbot, tmp_path, monkeypatch
    ) -> None:  # type: ignore[no-untyped-def]
        """J-1.1 step 21: sweep aperture 0.33 → 0.42 in 4 points, export the CSV."""
        import csv

        import numpy as np

        window = _window(qtbot)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            window.last_sweep_result = window.sensor.sweep(
                "optics.aperture_diameter_m", np.linspace(0.33, 0.42, 4), keep_results=True
            )
        window.action("file.export_sweep_csv").setEnabled(True)
        dest = tmp_path / "sweep.csv"
        _save_to(monkeypatch, dest)
        window.action("file.export_sweep_csv").trigger()
        body = [
            ln for ln in dest.read_text(encoding="utf-8").splitlines() if not ln.startswith("#")
        ]
        rows = list(csv.reader(body))
        header = rows[0]
        assert header[0] == "optics.aperture_diameter_m [m]"
        assert "snr" in header  # dimensionless: bare name
        assert "fwhm_x_m [m]" in header
        assert "niirs_extrapolated [0/1 flag]" in header  # F-31: self-describing
        assert "sampling_regime_code [code]" in header
        assert [row[0] for row in rows[1:]] == ["0.33", "0.36", "0.39", "0.42"]
        assert not any(cell.startswith("np.") for row in rows for cell in row)
        for row in rows[1:]:
            for cell in row:
                float(cell)  # every cell is a number


def _stamp_lines(path: Path) -> dict[str, str]:
    stamp: dict[str, str] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.startswith("# "):
            break
        key, _, value = line[2:].partition(": ")
        stamp[key] = value
    return stamp


class TestF34F35RunStamps:
    """F-34: a retained sweep exported after edits that made it stale, unmarked.
    F-35: after a failed re-evaluation the metrics export wrote the previous
    result with no stale flag, no run stamp."""

    def test_fresh_metrics_export_carries_a_run_stamp(self, qtbot, tmp_path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
        window = _window(qtbot)
        dest = tmp_path / "metrics.csv"
        _save_to(monkeypatch, dest)
        window.action("file.export_metrics_csv").trigger()
        stamp = _stamp_lines(dest)
        assert stamp["stale"] == "no"
        assert stamp["run_id"] and stamp["evaluated_at"] and stamp["radiant"].startswith("v")
        assert stamp["config"].endswith("mwir_leo_minimal.yaml")
        # The header follows the stamp lines unchanged.
        lines = dest.read_text(encoding="utf-8").splitlines()
        assert lines[len(stamp)] == "name,value,unit,description"

    def test_metrics_export_after_a_failed_reevaluation_is_marked_stale(
        self, qtbot, tmp_path, monkeypatch
    ) -> None:  # type: ignore[no-untyped-def]
        """T-G: success, then an edit that fails; the previous result is what exports."""
        from radiant.core.exceptions import CoreValidationError

        window = _window(qtbot)
        monkeypatch.setattr("radiant.gui.main_window.exec_dialog", lambda dlg, *a, **k: 0)
        window._on_eval_failed(CoreValidationError("some genuine rejection"))  # noqa: SLF001
        dest = tmp_path / "metrics.csv"
        _save_to(monkeypatch, dest)
        window.action("file.export_metrics_csv").trigger()
        assert _stamp_lines(dest)["stale"].startswith("yes — the last re-evaluation failed")
        assert "stale" in window.statusBar().currentMessage()

    def test_sweep_export_after_an_edit_is_marked_stale(self, qtbot, tmp_path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
        """T-F: sweep, then edit detector.qe_value, then File ▸ Export Sweep CSV."""
        import numpy as np

        window = _window(qtbot)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            sweep = window.sensor.sweep(
                "optics.aperture_diameter_m", np.linspace(0.25, 0.35, 3), keep_results=False
            )
        window.last_sweep_result = sweep
        window._sweep_run_at = "2026-09-20T00:00:00+00:00"  # noqa: SLF001 — as _on_run_sweep records
        window._sweep_model_serial = window._model_serial  # noqa: SLF001
        window.action("file.export_sweep_csv").setEnabled(True)
        fresh = tmp_path / "sweep_fresh.csv"
        _save_to(monkeypatch, fresh)
        window.action("file.export_sweep_csv").trigger()
        assert _stamp_lines(fresh)["stale"] == "no"

        window.sensor.set("detector.qe_value", 0.5)
        with qtbot.waitSignal(window.evaluationFinished, timeout=_WAIT_MS):
            window.parameter_panel.parameterEdited.emit("detector.qe_value")
        stale = tmp_path / "sweep_stale.csv"
        _save_to(monkeypatch, stale)
        window.action("file.export_sweep_csv").trigger()
        assert _stamp_lines(stale)["stale"].startswith("yes — the configuration was edited")

    def test_workbook_carries_a_run_sheet(self, qtbot, tmp_path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
        import openpyxl

        window = _window(qtbot)
        dest = tmp_path / "wb.xlsx"
        _save_to(monkeypatch, dest)
        window.action("file.export_xlsx").trigger()
        book = openpyxl.load_workbook(dest)
        assert "Run" in book.sheetnames
        rows = {row[0].value: row[1].value for row in book["Run"].iter_rows(min_row=2)}
        assert rows["stale"] == "no" and rows["run_id"]
